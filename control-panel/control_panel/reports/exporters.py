"""Report exports: Excel (every dataset, typed, with groups and totals), CSV
(one dataset) and PDF (the whole report, printable), all from one RunResult."""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle

from .engine import SUPPRESSED, DatasetResult, RunResult
from .model import INT, MONEY, NUMBER, PCT

_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")
_FORMATS = {INT: "#,##0", NUMBER: "#,##0.00", MONEY: "#,##0.00", PCT: "0.0%"}
_HEADER_FILL = PatternFill("solid", fgColor="2D2A5A")
_SUBTOTAL_FILL = PatternFill("solid", fgColor="E9E6F7")
_TOTAL_FILL = PatternFill("solid", fgColor="D6D0F0")


def _stamp() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")


def _filename(result: RunResult, ext: str, suffix: str = "") -> str:
    return f"report-{result.report.id}{'-' + suffix if suffix else ''}-{_stamp()}.{ext}"


def _safe_text(value) -> str:
    text = "" if value is None else str(value)
    return "'" + text if text.startswith(_FORMULA_PREFIX) else text


def _excel_value(cell, field):
    v = cell.value
    if v is None or v == SUPPRESSED:
        return SUPPRESSED if v == SUPPRESSED else None
    if field.type in (INT, NUMBER, MONEY, PCT) and isinstance(v, (int, float)):
        return v
    if isinstance(v, bool):
        return "Yes" if v else "No"
    return _safe_text(cell.text if isinstance(v, (list, dict)) else v)


def _parameters(result: RunResult) -> list[tuple[str, str]]:
    labels = {p.name: p.label for p in result.report.params}
    return [(labels.get(k, k), v) for k, v in result.params.items()]


def xlsx(result: RunResult) -> HttpResponse:
    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    for ds in result.datasets:
        name = ds.spec.title[:28] or ds.spec.key
        base, n = name, 2
        while name in used:
            name = f"{base[:25]} {n}"
            n += 1
        used.add(name)
        ws = wb.create_sheet(name.replace("/", "-").replace("\\", "-").replace("?", "").replace("*", "")
                             .replace("[", "(").replace("]", ")").replace(":", " "))
        ws.append([f.label for f in ds.fields])
        for i, f in enumerate(ds.fields, start=1):
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = _HEADER_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = max(f.width, min(len(f.label) + 2, 40))
        for row in ds.rows:
            ws.append([_excel_value(cell, f) for cell, f in zip(row.cells, ds.fields)])
            r = ws.max_row
            for i, f in enumerate(ds.fields, start=1):
                c = ws.cell(row=r, column=i)
                if f.type in _FORMATS and isinstance(c.value, (int, float)):
                    c.number_format = _FORMATS[f.type]
                if row.kind != "data":
                    c.font = Font(bold=True)
                    c.fill = _TOTAL_FILL if row.kind == "total" else _SUBTOTAL_FILL
        ws.freeze_panes = "A2"
        if ds.rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(ds.fields))}{ws.max_row}"
        if ds.spec.note or ds.truncated or ds.mixed_currency:
            ws.append([])
            for note in filter(None, (ds.spec.note,
                                      "Rows were capped; narrow the parameters for the rest." if ds.truncated else "",
                                      "Amounts are per currency; totals across currencies are not added." if ds.mixed_currency else "")):
                ws.append([note])
    about = wb.create_sheet("About", 0)
    about.append(["Report", result.report.title])
    about.append(["Category", result.report.category])
    about.append(["Description", result.report.description])
    about.append(["Generated (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])
    for label, value in _parameters(result):
        about.append([f"Parameter: {label}", value])
    about.append(["Source", f"/v1/admin/{result.report.admin_path}"])
    about.append(["Note", "Counts from 1 to 4 are shown as <5 (privacy)." ])
    about.column_dimensions["A"].width = 24
    about.column_dimensions["B"].width = 80
    for row in about.iter_rows(min_col=1, max_col=1):
        row[0].font = Font(bold=True)
    buffer = io.BytesIO()
    wb.save(buffer)
    return _attachment(buffer.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                       _filename(result, "xlsx"))


def csv_dataset(result: RunResult, ds: DatasetResult) -> HttpResponse:
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow([f.label for f in ds.fields])
    for row in ds.rows:
        values = []
        for cell, f in zip(row.cells, ds.fields):
            v = cell.value
            values.append(v if isinstance(v, (int, float)) and not isinstance(v, bool) else _safe_text(cell.text if v is not None else ""))
        writer.writerow(values)
    # BOM so Excel opens UTF-8 (₹, €) correctly.
    return _attachment(("﻿" + out.getvalue()).encode("utf-8"), "text/csv; charset=utf-8",
                       _filename(result, "csv", ds.spec.key))


def pdf(result: RunResult) -> HttpResponse:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=14 * mm, title=result.report.title,
                            author="Operations console")
    styles = getSampleStyleSheet()
    small = styles["BodyText"].clone("small", fontSize=7.5, leading=9)
    cell_style = styles["BodyText"].clone("cell", fontSize=7, leading=8.5)
    story = [Paragraph(result.report.title, styles["Title"]),
             Paragraph(result.report.description, small)]
    params = " · ".join(f"{label}: {value}" for label, value in _parameters(result)) or "Default parameters"
    story += [Paragraph(f"{params} · generated {datetime.now(timezone.utc):%Y-%m-%d %H:%M} UTC", small), Spacer(1, 6 * mm)]
    for err in result.errors:
        story.append(Paragraph(f"Unavailable: {err}", small))
    for ds in result.datasets:
        story.append(Paragraph(ds.spec.title, styles["Heading3"]))
        if not ds.rows:
            story += [Paragraph("No data for these parameters.", small), Spacer(1, 4 * mm)]
            continue
        data = [[Paragraph(f"<b>{f.label}</b>", cell_style) for f in ds.fields]]
        for row in ds.rows:
            data.append([Paragraph(_xml(c.text), cell_style) for c in row.cells])
        width = doc.width / max(1, len(ds.fields))
        table = LongTable(data, colWidths=[width] * len(ds.fields), repeatRows=1)
        style = [
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2D2A5A")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#B9B4D6")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F5FB")]),
        ]
        for i, row in enumerate(ds.rows, start=1):
            if row.kind != "data":
                style.append(("BACKGROUND", (0, i), (-1, i), colors.HexColor("#D6D0F0" if row.kind == "total" else "#E9E6F7")))
        table.setStyle(TableStyle(style))
        story += [table, Spacer(1, 3 * mm)]
        if ds.spec.note:
            story.append(Paragraph(ds.spec.note, small))
        story.append(Spacer(1, 5 * mm))

    def footer(canvas, doc_):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.drawString(12 * mm, 7 * mm, f"{result.report.title} · counts from 1 to 4 shown as <5")
        canvas.drawRightString(doc_.pagesize[0] - 12 * mm, 7 * mm, f"Page {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return _attachment(buffer.getvalue(), "application/pdf", _filename(result, "pdf"))


def _xml(text: str) -> str:
    return str(text).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _attachment(body: bytes, content_type: str, filename: str) -> HttpResponse:
    response = HttpResponse(body, content_type=content_type)
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    response["Cache-Control"] = "no-store"
    return response
