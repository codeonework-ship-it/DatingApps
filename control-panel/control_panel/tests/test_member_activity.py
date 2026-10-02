"""Member activity: the explorer, the member page section, and the live tail."""
from __future__ import annotations

import io
from unittest.mock import patch

from asgiref.sync import sync_to_async
from django.urls import reverse
from openpyxl import load_workbook

from control_panel import consumers
from control_panel.services.go_client import APIResult

from .case_support import AutoOKClient, ConsoleCaseTest, bff_error
from .test_live import LiveSocketBase

MEMBER = "00000000-0000-0000-0000-0000000000c1"
ACTION = {
    "id": "req:42", "at": "2026-10-02T10:15:00Z", "source": "request", "member_id": MEMBER, "actor_id": MEMBER,
    "actor_role": "member", "category": "Safety", "action_key": "safety.report", "action_label": "Reported a member",
    "method": "POST", "route": "/v1/reports", "status_code": 201, "outcome": "success", "duration_ms": 41,
    "entity_type": "report", "entity_id": "r-9", "ip": "203.0.113.7", "device_id": "dev-1", "platform": "android",
    "app_version": "1.4.0", "user_agent": "<script>ua</script>", "request_id": "req-1", "correlation_id": "corr-1",
    "session_id": "sess-1", "details": {"reason": "spam"},
}


class MemberActivityPageTest(ConsoleCaseTest):
    module = "control_panel.views_activity"

    def test_lists_actions_with_detail_and_escapes_text(self):
        """Each action shows who, what, outcome, device and IP, with a detail row. [case:console.activity.explorer.renders]"""
        api = self.bff()
        api.list_member_actions.return_value = APIResult(True, {"actions": [ACTION], "total": 1})
        response = self.client.get(reverse("member_activity"))
        self.assertEqual(response.status_code, 200)
        for text in ("Reported a member", "safety.report", "203.0.113.7", "android 1.4.0", "corr-1", "reason: spam"):
            self.assertContains(response, text)
        self.assertContains(response, "&lt;script&gt;ua&lt;/script&gt;")
        self.assertNotContains(response, "<script>ua</script>")
        self.assertContains(response, reverse("user_detail", args=[MEMBER]))
        self.assertContains(response, 'data-live-topic="activity"')

    def test_filters_reach_go(self):
        """Member, area, outcome, source, reads and dates are server-side filters. [case:console.activity.explorer.filters]"""
        api = self.bff()
        api.list_member_actions.return_value = APIResult(True, {"actions": [], "total": 0})
        self.client.get(reverse("member_activity"), {"member": MEMBER, "category": "Safety", "outcome": "client_error",
                                                     "source": "request", "include_reads": "true", "from": "2026-10-01",
                                                     "q": "report", "category_bad": "x", "page": "2"})
        api.list_member_actions.assert_called_once_with(
            limit=50, offset=50, member=MEMBER, category="Safety", outcome="client_error", source="request",
            include_reads="true", **{"from": "2026-10-01"}, q="report")

    def test_export_has_every_detail_column(self):
        """Excel carries IP, device, session, request and correlation ids. [case:console.activity.explorer.export]"""
        api = self.bff()
        api.list_member_actions.return_value = APIResult(True, {"actions": [ACTION], "total": 1})
        response = self.client.get(reverse("member_activity"), {"export": "xlsx"})
        rows = list(load_workbook(io.BytesIO(response.content)).worksheets[0].iter_rows(values_only=True))
        header = rows[0]
        for column in ("IP address", "Device ID", "Session ID", "Request ID", "Correlation ID", "User agent", "Details"):
            self.assertIn(column, header)
        self.assertEqual(rows[1][header.index("IP address")], "203.0.113.7")
        self.assertEqual(rows[1][header.index("Details")], "reason: spam")

    def test_failure_shows_banner(self):
        """[case:console.activity.explorer.failure]"""
        self.bff().list_member_actions.return_value = bff_error()
        self.assertContains(self.client.get(reverse("member_activity")), "alert-glass warning")


class MemberPageActivityTest(ConsoleCaseTest):
    def test_member_page_shows_summary_and_latest_actions(self):
        """A member's page shows their latest actions and links to the full log. [case:console.activity.member.section]"""
        api = self.bff()
        api.get_user.return_value = APIResult(True, {"user": {"id": MEMBER, "username": "asha"}})
        api.member_activity.return_value = APIResult(True, {"actions": [ACTION], "total": 120, "summary": {
            "first_seen": "2026-09-01T08:00:00Z", "last_seen": "2026-10-02T10:15:00Z", "devices": 2, "ips": 3,
            "by_category": [{"category": "Safety", "count": 4}]}})
        response = self.client.get(reverse("user_detail", args=[MEMBER]))
        self.assertContains(response, "Reported a member")
        self.assertContains(response, "Showing the latest 1 of 120 actions.")
        self.assertContains(response, "Safety · 4")
        self.assertContains(response, reverse("member_activity") + "?member=" + MEMBER)
        api.member_activity.assert_called_once_with(MEMBER, limit=25)

    def test_role_without_access_sees_a_note(self):
        """[case:console.activity.member.denied]"""
        api = self.bff()
        api.get_user.return_value = APIResult(True, {"user": {"id": MEMBER, "username": "asha"}})
        api.member_activity.return_value = bff_error("forbidden", status=403)
        self.assertContains(self.client.get(reverse("user_detail", args=[MEMBER])), "not available to your role")


class ActivityLiveTailTest(LiveSocketBase):
    async def test_tail_primes_then_pushes_new_rows_for_the_page_filters(self):
        """The live tail sends only actions newer than the page, with its filters. [case:console.activity.live_tail]"""
        self.api.activity_stream.side_effect = [
            APIResult(True, {"actions": [ACTION], "cursor": "c1"}),
            APIResult(True, {"actions": [dict(ACTION, id="req:43", action_label="Sent a message")], "cursor": "c2"}),
        ] + [APIResult(True, {"actions": [], "cursor": "c2"})] * 20
        _, key = await sync_to_async(self._session)()
        socket = self._socket(key)
        connected, _ = await socket.connect()
        self.assertTrue(connected)
        await socket.send_json_to({"type": "subscribe", "topics": ["activity"],
                                   "activity": {"member": MEMBER, "category": "Safety", "source": "nope", "evil": "x"}})
        message = await socket.receive_json_from(timeout=10)
        self.assertEqual(message["topic"], "activity")
        self.assertEqual(len(message["data"]["rows"]), 1)
        self.assertIn("Sent a message", message["data"]["rows"][0])
        self.assertNotIn("Reported a member", message["data"]["rows"][0])  # primed, not re-sent
        first, second = self.api.activity_stream.call_args_list[:2]
        self.assertEqual(first.kwargs, {"limit": 1, "member": MEMBER, "category": "Safety"})
        self.assertEqual(second.kwargs, {"after": "c1", "limit": 100, "member": MEMBER, "category": "Safety"})
        await socket.disconnect()
