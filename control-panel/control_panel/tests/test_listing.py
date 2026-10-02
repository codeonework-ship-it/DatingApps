"""Server-side lists (control_panel/listing.py): paging, search and filters
are read safely, Go does the work, and Excel exports the whole result."""
from __future__ import annotations

import io

from django.urls import reverse
from openpyxl import load_workbook

from control_panel.services.go_client import APIResult

from .case_support import ConsoleCaseTest

MEMBER = "00000000-0000-0000-0000-0000000000b1"


def _users(n: int, start: int = 0) -> list[dict]:
    return [{"id": f"u{start + i}", "name": f"Member {start + i}", "username": f"m{start + i}",
             "created_at": "2026-09-01T10:00:00Z", "is_verified": i % 2 == 0} for i in range(n)]


class UserListPagingTest(ConsoleCaseTest):
    def test_page_size_and_page_reach_go_as_limit_and_offset(self):
        """Paging is server side. [case:console.lists.paging.server_side]"""
        api = self.bff()
        api.list_users.return_value = APIResult(True, {"users": _users(10), "total": 95})
        response = self.client.get(reverse("user_list"), {"page": "4", "page_size": "10", "gender": "female"})
        self.assertEqual(response.status_code, 200)
        api.list_users.assert_called_once_with(limit=10, offset=30, q="", status="", gender="female", verified="")
        self.assertContains(response, "Showing <strong>31–40</strong> of <strong>95</strong>")
        page = response.context["page"]
        self.assertEqual(page.pages, 10)
        self.assertIn("page=5", page.next_url)
        self.assertIn("gender=female", page.next_url)
        self.assertIn("page_size=10", page.next_url)

    def test_unknown_values_never_reach_go(self):
        """Filters, sizes and pages outside the allow-list are dropped. [case:console.lists.params.allow_list]"""
        api = self.bff()
        api.list_users.return_value = APIResult(True, {"users": [], "total": 0})
        self.client.get(reverse("user_list"), {"status": "'; drop table", "gender": "robot", "page_size": "100000",
                                               "page": "-4", "verified": "yes"})
        api.list_users.assert_called_once_with(limit=25, offset=0, q="", status="", gender="", verified="yes")

    def test_search_text_is_escaped_in_page_links(self):
        """A search with & or quotes keeps its meaning in every link. [case:console.lists.search.escaped]"""
        api = self.bff()
        api.list_users.return_value = APIResult(True, {"users": _users(25), "total": 60})
        response = self.client.get(reverse("user_list"), {"q": 'a&status=banned"x'})
        api.list_users.assert_called_once_with(limit=25, offset=0, q='a&status=banned"x', status="", gender="", verified="")
        self.assertIn("q=a%26status%3Dbanned%22x", response.context["page"].next_url)
        self.assertNotContains(response, 'banned"x"')

    def test_a_page_past_the_end_shows_the_last_page(self):
        """Filtering down while on page 9 lands on the last page. [case:console.lists.paging.clamps]"""
        api = self.bff()
        api.list_users.side_effect = [APIResult(True, {"users": [], "total": 30}), APIResult(True, {"users": _users(5, 25), "total": 30})]
        response = self.client.get(reverse("user_list"), {"page": "9"})
        self.assertEqual(response.context["page"].query.page, 2)
        self.assertEqual(api.list_users.call_args.kwargs["offset"], 25)


class ExcelExportTest(ConsoleCaseTest):
    def test_export_pages_through_go_and_returns_a_workbook(self):
        """Export Excel holds every filtered row, not just one page. [case:console.lists.export.xlsx]"""
        api = self.bff()
        api.list_users.side_effect = [
            APIResult(True, {"users": _users(200), "total": 250}),
            APIResult(True, {"users": _users(50, 200), "total": 250}),
        ]
        response = self.client.get(reverse("user_list"), {"export": "xlsx", "status": "active", "page": "3"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn('filename="users-', response["Content-Disposition"])
        self.assertEqual([c.kwargs["offset"] for c in api.list_users.call_args_list], [0, 200])
        self.assertTrue(all(c.kwargs["status"] == "active" for c in api.list_users.call_args_list))
        wb = load_workbook(io.BytesIO(response.content))
        sheet = wb[wb.sheetnames[0]]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][:3], ("User ID", "Name", "Username"))
        self.assertEqual(len(rows), 251)
        about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True)}
        self.assertEqual(about["Rows"], 250)
        self.assertEqual(about["Filter: status"], "active")

    def test_cells_are_never_formulas(self):
        """Member text starting with = is written as text (no Excel injection). [case:console.lists.export.no_formulas]"""
        api = self.bff()
        api.list_users.return_value = APIResult(True, {"users": [{"id": MEMBER, "name": '=HYPERLINK("http://x","y")'}], "total": 1})
        response = self.client.get(reverse("user_list"), {"export": "xlsx"})
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        cell = sheet.cell(row=2, column=2)
        self.assertEqual(cell.data_type, "s")
        self.assertTrue(cell.value.startswith("'="))

    def test_export_failure_is_a_clear_error(self):
        """Go failing the first read gives a 502 with the reason. [case:console.lists.export.failure]"""
        from .case_support import bff_error
        self.bff().list_users.return_value = bff_error()
        response = self.client.get(reverse("user_list"), {"export": "xlsx"})
        self.assertEqual(response.status_code, 502)
        self.assertIn(b"Export failed", response.content)


class BillingListContractTest(ConsoleCaseTest):
    def test_search_and_dates_reach_go_and_export_matches(self):
        """Billing lists pass q and the date range to Go; Excel uses the same filters. [case:console.lists.contract.billing]"""
        api = self.bff()
        api.list_payments.return_value = APIResult(True, {"total": 1, "payments": [
            {"user_id": MEMBER, "amount_paise": 19900, "currency": "INR", "status": "captured"}]})
        self.client.get(reverse("billing_payments"), {"q": "pay_1", "from": "2026-09-01", "to": "2026-09-30", "status": "captured"})
        api.list_payments.assert_called_once_with(limit=25, offset=0, status="captured", **{"from": "2026-09-01", "to": "2026-09-30", "q": "pay_1"})
        response = self.client.get(reverse("billing_payments"), {"export": "xlsx", "from": "2026-09-01", "to": "bad"})
        self.assertEqual(api.list_payments.call_args.kwargs, {"limit": 200, "offset": 0, "from": "2026-09-01"})
        sheet = load_workbook(io.BytesIO(response.content)).worksheets[0]
        self.assertEqual(sheet.cell(row=2, column=2).value, 199.0)
