"""Report server: parameters, groups and totals, suppression, drill-through,
roles, and the Excel, CSV and PDF exports."""
from __future__ import annotations

import csv
import io

from django.urls import reverse
from openpyxl import load_workbook

from control_panel.services.go_client import APIResult

from .case_support import ConsoleCaseTest, bff_error

REVENUE = {
    "totals": [
        {"currency": "INR", "market": "IN", "gross": "1000.00", "gross_minor": 100000, "net": "900.00", "net_minor": 90000,
         "paying_members": 12, "refund_rate": 0.05},
        # Go writes the suppression marker into the value itself (suppressCount, business_reports.go).
        {"currency": "INR", "market": "IN-KA", "gross": "500.00", "gross_minor": 50000, "net": "450.00", "net_minor": 45000,
         "paying_members": "<5", "refund_rate": 0.0},
        {"currency": "EUR", "market": "DE", "gross": "20.00", "gross_minor": 2000, "net": "20.00", "net_minor": 2000,
         "paying_members": 7, "refund_rate": 0.1},
    ],
    "by_city": [{"city": "Bengaluru", "currency": "INR", "paying_members": 9, "net": "300.00", "net_minor": 30000}],
    "trend": [{"bucket": "2026-09-28", "currency": "INR", "net_minor": 90000, "refunded_minor": 1000}],
}


class ReportServerTest(ConsoleCaseTest):
    module = "control_panel.reports.engine"

    def _revenue(self, **query):
        api = self.bff()
        api.business_report.return_value = APIResult(True, REVENUE)
        return api, self.client.get(reverse("report_view", args=["revenue"]), query)

    def test_catalog_lists_reports_and_searches(self):
        """The catalog lists business and product reports and finds them by keyword. [case:console.reports.catalog.search] [case:console.reports.report_catalog.renders] [case:console.reports.report_catalog.filters]"""
        self.bff()
        response = self.client.get(reverse("report_catalog"))
        card = '<span class="report-card-title">'
        self.assertContains(response, card + "Revenue<")
        self.assertContains(response, card + "Retention cohorts<")
        found = self.client.get(reverse("report_catalog"), {"q": "churn"})
        self.assertContains(found, card + "Subscriptions &amp; MRR<")
        self.assertNotContains(found, card + "Coin economy<")

    def test_catalog_hides_reports_the_role_cannot_read(self):
        """A support agent sees no finance reports. [case:console.reports.catalog.role_scoped]"""
        self.bff()
        response = self.as_roles("support").get(reverse("report_catalog"))
        self.assertNotContains(response, '<span class="report-card-title">Revenue<')

    def test_parameters_are_validated_before_reaching_go(self):
        """Only listed parameter values reach Go; defaults fill the rest. [case:console.reports.params.validated]"""
        api, _ = self._revenue(since="2026-09-01", until="not-a-date", tz="Mars/Base", mode="sandbox", bucket="month", evil="1")
        api.business_report.assert_called_once_with("revenue", {"since": "2026-09-01", "tz": "UTC", "mode": "sandbox", "bucket": "month"})

    def test_groups_subtotals_and_mixed_currency_totals(self):
        """Grouped rows get subtotals; money is never added across currencies. [case:console.reports.view.groups_totals]"""
        _, response = self._revenue()
        self.assertEqual(response.status_code, 200)
        totals = response.context["result"].datasets[0]
        self.assertEqual(totals.group_by, ("currency",))
        self.assertEqual([g.label for g in totals.groups], ["INR", "EUR"])
        inr = totals.groups[0].subtotal
        gross = [c for c, f in zip(inr.cells, totals.fields) if f.key == "gross"][0]
        self.assertEqual(gross.value, 1500.0)
        self.assertEqual(gross.text, "1,500.00 INR")
        grand = [c for c, f in zip(totals.total.cells, totals.fields) if f.key == "gross"][0]
        self.assertEqual(grand.text, "mixed currencies")
        self.assertContains(response, "Totals across currencies are not added.")
        self.assertContains(response, "&lt;5")  # suppressed small count stays suppressed

    def test_group_by_can_be_changed_or_removed(self):
        """Operators regroup a table or show it flat. [case:console.reports.view.group_by]"""
        _, by_market = self._revenue(group_totals="market")
        self.assertEqual(by_market.context["result"].datasets[0].group_by, ("market",))
        _, flat = self._revenue(group_totals="")
        self.assertEqual(flat.context["result"].datasets[0].groups, [])
        _, refused = self._revenue(group_totals="gross")  # not a groupable field
        self.assertEqual(refused.context["result"].datasets[0].group_by, ())

    def test_city_drills_through_to_liquidity(self):
        """A city links to the liquidity report for that city. [case:console.reports.view.drill_through]"""
        _, response = self._revenue(since="2026-09-01")
        self.assertContains(response, reverse("report_view", args=["liquidity"]) + "?since=2026-09-01&amp;city=Bengaluru")

    def test_role_refusal_is_shown_not_crashed(self):
        """Go's 403 renders as not available to your role. [case:console.reports.view.denied] [case:console.reports.report_view.renders]"""
        self.bff().business_report.return_value = bff_error("forbidden", status=403)
        response = self.client.get(reverse("report_view", args=["revenue"]))
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "not available to your role", status_code=403)

    def test_unknown_report_is_404(self):
        """[case:console.reports.view.unknown] [case:console.reports.report_view.renders]"""
        self.bff()
        self.assertEqual(self.client.get(reverse("report_view", args=["nope"])).status_code, 404)

    def test_excel_export_has_typed_sheets_and_parameters(self):
        """Excel: one typed sheet per table, subtotal rows, and an About sheet. [case:console.reports.export.xlsx]"""
        _, response = self._revenue(export="xlsx", mode="all")
        self.assertEqual(response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(io.BytesIO(response.content))
        self.assertEqual(wb.sheetnames[0], "About")
        about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True)}
        self.assertEqual(about["Parameter: Payments"], "all")
        sheet = wb[wb.sheetnames[1]]
        rows = list(sheet.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Currency")
        gross_col = rows[0].index("Gross")
        self.assertEqual(rows[1][gross_col], 1000.0)  # a number, not text
        self.assertIn("INR subtotal", [r[0] for r in rows])

    def test_csv_export_of_one_table(self):
        """CSV exports one table, UTF-8 with BOM for Excel. [case:console.reports.export.csv]"""
        _, response = self._revenue(export="csv", dataset="by_city")
        self.assertTrue(response.content.startswith("﻿".encode()))
        rows = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["City", "Currency", "Paying members", "Net"])
        self.assertEqual(rows[1][0], "Bengaluru")

    def test_pdf_export(self):
        """PDF export renders the whole report. [case:console.reports.export.pdf]"""
        _, response = self._revenue(export="pdf")
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_analytics_reports_use_go_column_metadata(self):
        """Product reports take their tables and units from Go. [case:console.reports.view.analytics_tables] [case:console.reports.report_view.renders]"""
        api = self.bff()
        api.analytics_report.return_value = APIResult(True, {"tables": {"triangle": {
            "columns": [{"key": "cohort_week", "label": "Week", "kind": "dimension"},
                        {"key": "week_1", "label": "Week 1", "kind": "ratio", "unit": "%"}],
            "rows": [{"cohort_week": "2026-09-01", "week_1": 42.5, "suppressed": []},
                     {"cohort_week": "2026-09-08", "week_1": None, "suppressed": ["week_1"]}]}}})
        response = self.client.get(reverse("report_view", args=["retention"]), {"grain": "week"})
        self.assertContains(response, "Week 1 (%)")
        self.assertContains(response, "42.50")
        self.assertContains(response, "&lt;5")
        api.analytics_report.assert_called_once_with("retention", {"grain": "week"})


MEMBER_ID = "00000000-0000-0000-0000-0000000000d4"


class MemberReportsTest(ConsoleCaseTest):
    module = "control_panel.reports.engine"

    def test_member_360_asks_for_a_member_first(self):
        """No member, no Go calls: the report asks for one. [case:console.reports.member_360.requires_member]"""
        api = self.bff()
        response = self.client.get(reverse("report_view", args=["member-360"]))
        self.assertContains(response, "Enter Member ID or @username")
        api.get_user.assert_not_called()

    def test_member_360_resolves_a_username_and_reads_every_section(self):
        """@username resolves to the member; each section asks Go for that member only. [case:console.reports.member_360.renders]"""
        api = self.bff()
        api.list_users.return_value = APIResult(True, {"users": [{"id": MEMBER_ID, "username": "asha"}]})
        api.get_user.return_value = APIResult(True, {"user": {"id": MEMBER_ID, "username": "asha", "is_verified": True,
                                                              "profile_completion": 80}})
        api.get_wallet_balance.return_value = APIResult(True, {"wallet": {"coin_balance": 250}})
        api.member_activity.return_value = APIResult(True, {"summary": {"total": 12, "distinct_devices": 2,
                                                                        "by_category": {"Safety": 3, "Auth": 9, "Billing & coins": 0}},
                                                            "actions": [{"at": "2026-10-02T10:00:00Z", "action_label": "Signed in"}]})
        api.list_reports.return_value = APIResult(True, {"reports": [{"reported_user_id": "x", "reason": "spam"}], "total": 1})
        response = self.client.get(reverse("report_view", args=["member-360"]), {"member": "@asha", "from": "2026-09-01"})
        self.assertEqual(response.status_code, 200)
        api.get_user.assert_called_once_with(MEMBER_ID)
        self.assertEqual(api.list_reports.call_args_list[0].kwargs, {"limit": 500, "offset": 0, "reporter_user_id": MEMBER_ID, "from": "2026-09-01"})
        self.assertEqual(api.list_reports.call_args_list[1].kwargs["reported_user_id"], MEMBER_ID)
        self.assertEqual(api.list_support_tickets.call_args.kwargs["member"], MEMBER_ID)
        self.assertEqual(api.list_payments.call_args.kwargs["user_id"], MEMBER_ID)
        result = response.context["result"]
        by_key = {d.spec.key: d for d in result.datasets}
        profile = {r["label"]: r["value"] for r in by_key["profile"].raw_rows}
        self.assertEqual(profile["Verified"], "Yes")
        self.assertEqual(profile["Profile complete (%)"], "80")
        self.assertEqual([r["area"] for r in by_key["activity_by_area"].raw_rows], ["Auth", "Safety"])
        self.assertContains(response, "Signed in")
        self.assertContains(response, "250")

    def test_unknown_username_is_explained(self):
        """[case:console.reports.member_360.unknown_member]"""
        self.bff().list_users.return_value = APIResult(True, {"users": [{"id": "z", "username": "ashanti"}]})
        response = self.client.get(reverse("report_view", args=["member-360"]), {"member": "asha"})
        self.assertContains(response, "No member with username @asha.")

    def test_directory_pages_through_go_and_drills_into_member_360(self):
        """The directory reads every page and links each member to Member 360. [case:console.reports.member_directory.renders]"""
        api = self.bff()
        api.list_users.side_effect = [
            APIResult(True, {"users": [{"id": f"u{i}", "username": f"m{i}", "city": "Pune"} for i in range(500)], "total": 501}),
            APIResult(True, {"users": [{"id": "u500", "username": "m500", "city": "Goa"}], "total": 501}),
        ]
        response = self.client.get(reverse("report_view", args=["member-directory"]), {"verified": "yes"})
        self.assertEqual([c.kwargs["offset"] for c in api.list_users.call_args_list], [0, 500])
        self.assertEqual(api.list_users.call_args.kwargs["verified"], "yes")
        self.assertEqual(len(response.context["result"].datasets[0].raw_rows), 501)
        self.assertContains(response, reverse("report_view", args=["member-360"]) + "?member=u500")

    def test_most_reported_lists_the_largest_groups_first(self):
        """Repeat offenders come first. [case:console.reports.most_reported_members.renders]"""
        api = self.bff()
        api.list_reports.return_value = APIResult(True, {"reports": [
            {"reported_user_id": "a"}, {"reported_user_id": "b"}, {"reported_user_id": "b"}, {"reported_user_id": "b"}], "total": 4})
        response = self.client.get(reverse("report_view", args=["most-reported-members"]))
        groups = response.context["result"].datasets[0].groups
        self.assertEqual([(g.label, g.count) for g in groups], [("b", 3), ("a", 1)])


class OperationsReportsTest(ConsoleCaseTest):
    module = "control_panel.reports.engine"

    def test_queue_sla_snapshot_and_oldest_first(self):
        """Each queue's open total and oldest age; open items oldest first with past-target flags. [case:console.reports.queue_sla.renders]"""
        api = self.bff()
        api.list_reports.side_effect = lambda **kw: APIResult(True, {
            "reports": [{"created_at": "2020-01-01T00:00:00Z", "status": "pending"}] if kw.get("limit") == 1 else [
                {"created_at": "2099-01-01T00:00:00Z", "status": "pending", "reason": "new"},
                {"created_at": "2020-01-01T00:00:00Z", "status": "pending", "reason": "old"}],
            "total": 2})
        response = self.client.get(reverse("report_view", args=["queue-sla"]))
        by_key = {d.spec.key: d for d in response.context["result"].datasets}
        snapshot = {r["queue"]: r for r in by_key["queues"].raw_rows}
        self.assertEqual(snapshot["Reports"]["open"], 2)
        self.assertTrue(snapshot["Reports"]["past_target"])
        self.assertEqual([r["reason"] for r in by_key["reports_open"].raw_rows], ["old", "new"])
        self.assertEqual(api.list_reports.call_args_list[0].kwargs, {"limit": 1, "offset": 0, "order": "asc", "status": "pending"})

    def test_dormant_members_filters_and_sorts_by_inactivity(self):
        """Only members inactive at least N days, longest first. [case:console.reports.dormant_members.renders] [case:console.reports.dormant_members.filters]"""
        api = self.bff()
        api.list_users.return_value = APIResult(True, {"users": [
            {"id": "a", "username": "recent", "last_login_at": "2099-01-01T00:00:00Z"},
            {"id": "b", "username": "old", "last_login_at": "2020-01-01T00:00:00Z"},
            {"id": "c", "username": "older", "last_login_at": "2019-01-01T00:00:00Z"}], "total": 3})
        response = self.client.get(reverse("report_view", args=["dormant-members"]), {"days": "30"})
        rows = response.context["result"].datasets[0].raw_rows
        self.assertEqual([r["username"] for r in rows], ["older", "old"])
        self.assertEqual(api.list_users.call_args.kwargs["status"], "active")

    def test_signin_security_reads_auth_actions(self):
        """Account security reads member Auth actions only. [case:console.reports.signin_security.renders]"""
        api = self.bff()
        api.list_member_actions.return_value = APIResult(True, {"actions": [{"action_label": "Signed in", "ip": "203.0.113.9"}], "total": 1})
        response = self.client.get(reverse("report_view", args=["signin-security"]), {"from": "2026-10-01"})
        kwargs = api.list_member_actions.call_args.kwargs
        self.assertEqual((kwargs["category"], kwargs["source"], kwargs["from"]), ("Auth", "request", "2026-10-01"))
        self.assertContains(response, "203.0.113.9")

    def test_daily_operations_survives_unavailable_sources(self):
        """A failing section is reported, the rest still renders, and it exports to PDF. [case:console.reports.daily_operations.renders] [case:console.reports.daily_operations.export_pdf]"""
        api = self.bff()
        api.system_requests.return_value = bff_error("not deployed", status=404)
        api.system_jobs.return_value = bff_error("not deployed", status=404)
        response = self.client.get(reverse("report_view", args=["daily-operations"]))
        self.assertEqual(response.status_code, 200)
        by_key = {d.spec.key: d for d in response.context["result"].datasets}
        self.assertEqual(by_key["server"].raw_rows, [{"signal": "Server activity", "value": "unavailable"}])
        pdf = self.client.get(reverse("report_view", args=["daily-operations"]), {"export": "pdf"})
        self.assertTrue(pdf.content.startswith(b"%PDF"))

    def test_sos_minutes_to_resolve(self):
        """[case:console.reports.sos_incidents.renders]"""
        api = self.bff()
        api.list_sos_alerts.return_value = APIResult(True, {"alerts": [
            {"triggered_at": "2026-10-01T10:00:00Z", "resolved_at": "2026-10-01T10:12:30Z", "status": "resolved"}], "total": 1})
        response = self.client.get(reverse("report_view", args=["sos-incidents"]))
        row = response.context["result"].datasets[0].rows[0]
        minutes = [c.value for c, f in zip(row.cells, response.context["result"].datasets[0].fields) if f.key == "minutes_to_resolve"][0]
        self.assertEqual(minutes, 12.5)
