"""Server-side paging for the remaining command-center lists: gift catalog,
daily prompts, match nudges, group covers, blog review cases, photo themes,
growth fraud signals, coin economy fraud cases and the XP fraud queue.

Every fixture below is shaped like the Go handler's real answer
(backend/internal/bff/mobile): the same items key and row keys, plus the
paging contract's total, limit and offset.
"""
from __future__ import annotations

import io
from unittest.mock import patch

from django.urls import reverse
from openpyxl import load_workbook

from control_panel.services.go_client import APIResult, GoBFFClient
from control_panel.tests.case_support import ConsoleCaseTest, bff_error

MEMBER = "11111111-1111-4111-8111-111111111111"
OTHER = "22222222-2222-4222-8222-222222222222"
CASE = "33333333-3333-4333-8333-333333333333"


def _page(items_key: str, rows: list[dict], total: int, *, limit: int = 25, offset: int = 0, **extra) -> APIResult:
    return APIResult(True, {items_key: rows, "total": total, "limit": limit, "offset": offset, **extra})


def _sheet(response) -> tuple[list[tuple], dict]:
    wb = load_workbook(io.BytesIO(response.content))
    rows = list(wb[wb.sheetnames[0]].iter_rows(values_only=True))
    about = {r[0]: r[1] for r in wb["About"].iter_rows(values_only=True)}
    return rows, about


# listGiftCatalogPage (admin_list_endpoints.go): gifts, count, source + paging.
def _gift(i: int, **overrides) -> dict:
    gift = {"id": f"gift_{i:03d}", "name": f"Gift {i}", "gif_url": "", "tier": "rare", "price_coins": 10 + i,
            "icon_key": "rose", "icon_emoji": "🌹", "category": "roses", "description": f"Gift number {i}",
            "max_per_match_per_day": 3, "is_active": i % 2 == 0, "sort_order": i, "start_date": None, "end_date": None,
            "created_at": "2026-09-01T10:00:00Z", "updated_at": "2026-09-02T10:00:00Z"}
    gift.update(overrides)
    return gift


def _catalog(rows, total, **kw):
    return _page("gifts", rows, total, count=len(rows), source="db", **kw)


class CatalogPagingTest(ConsoleCaseTest):
    def test_page_filters_and_sort_reach_go(self):
        """Page 3 of 10 rows is offset 20; category, tier, active, search and sort go to Go as the contract names them. [case:console.catalog.catalog_list.renders] [case:console.lists.paging.server_side]"""
        api = self.bff()
        api.list_catalog_gifts.return_value = _catalog([_gift(i) for i in range(20, 30)], 47, limit=10, offset=20)
        response = self.client.get(reverse("catalog_list"), {"page": "3", "page_size": "10", "category": "roses", "tier": "rare",
                                                             "active": "yes", "q": "gift", "sort": "price_coins", "dir": "desc"})
        self.assertEqual(response.status_code, 200)
        api.list_catalog_gifts.assert_called_once_with(limit=10, offset=20, category="roses", tier="rare", active="yes",
                                                       q="gift", sort="price_coins", order="desc")
        self.assertContains(response, "Showing <strong>21–30</strong> of <strong>47</strong>")
        self.assertContains(response, "Page 3 of 5")
        next_url = response.context["page"].next_url
        for part in ("page=4", "page_size=10", "category=roses", "dir=desc", "sort=price_coins"):
            self.assertIn(part, next_url)
        # Every row keeps its actions.
        self.assertContains(response, reverse("catalog_edit", args=["gift_020"]))
        self.assertContains(response, reverse("catalog_toggle", args=["gift_020"]))
        self.assertContains(response, 'data-confirm-title="Delete this gift?"')
        # The hand-built offset links are gone.
        self.assertNotContains(response, "offset=")

    def test_display_order_is_ascending_by_default_and_descending_is_kept(self):
        """The catalog reads in display order (ascending) by default; choosing descending survives paging. [case:console.catalog.catalog_list.renders]"""
        api = self.bff()
        api.list_catalog_gifts.return_value = _catalog([_gift(i) for i in range(25)], 60)
        response = self.client.get(reverse("catalog_list"))
        self.assertEqual(api.list_catalog_gifts.call_args.kwargs["order"], "asc")
        self.assertEqual(response.context["page"].next_url, "?page=2")
        self.assertContains(response, '<option value="asc" selected>Ascending</option>', html=False)
        response = self.client.get(reverse("catalog_list"), {"dir": "desc"})
        self.assertEqual(api.list_catalog_gifts.call_args.kwargs["order"], "desc")
        self.assertIn("dir=desc", response.context["page"].next_url)

    def test_unknown_filter_values_are_dropped(self):
        """A category, tier or status outside the console's lists never reaches Go. [case:console.lists.params.allow_list]"""
        api = self.bff()
        api.list_catalog_gifts.return_value = _catalog([], 0)
        self.client.get(reverse("catalog_list"), {"category": "weapons", "tier": "mythic", "active": "maybe", "sort": "price; drop"})
        api.list_catalog_gifts.assert_called_once_with(limit=25, offset=0, sort="sort_order", order="asc")

    def test_excel_export_holds_every_matching_gift(self):
        """Export Excel pages through Go with the same filters and writes Go's keys. [case:console.lists.export.xlsx] [case:console.catalog.catalog_list.export]"""
        api = self.bff()
        api.list_catalog_gifts.side_effect = [_catalog([_gift(i) for i in range(200)], 230, limit=200),
                                              _catalog([_gift(i) for i in range(200, 230)], 230, limit=200, offset=200)]
        response = self.client.get(reverse("catalog_list"), {"export": "xlsx", "category": "roses", "page": "4"})
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="gift-catalog-', response["Content-Disposition"])
        self.assertEqual([(c.kwargs["limit"], c.kwargs.get("offset")) for c in api.list_catalog_gifts.call_args_list], [(200, 0), (200, 200)])
        self.assertTrue(all(c.kwargs["category"] == "roses" for c in api.list_catalog_gifts.call_args_list))
        rows, about = _sheet(response)
        self.assertEqual(len(rows), 231)
        header = rows[0]
        self.assertEqual(header[:5], ("Gift ID", "Name", "Category", "Tier", "Coin cost"))
        first = dict(zip(header, rows[1]))
        self.assertEqual(first["Gift ID"], "gift_000")
        self.assertEqual(first["Coin cost"], 10)
        self.assertEqual(first["Active"], "Yes")
        self.assertEqual(first["Max per match per day"], 3)
        self.assertEqual(about["Rows"], 230)
        self.assertEqual(about["Filter: category"], "roses")

    def test_gift_edit_still_finds_gifts_on_later_pages(self):
        """The edit page pages through Go's 500-row pages to find its gift. [case:console.catalog.catalog_edit.performs]"""
        api = self.bff()
        api.list_catalog_gifts.side_effect = [_catalog([_gift(i) for i in range(500)], 501, limit=500),
                                              _catalog([_gift(500, name="Late gift")], 501, limit=500, offset=500)]
        response = self.client.get(reverse("catalog_edit", args=["gift_500"]))
        self.assertContains(response, 'value="Late gift"')


# listEngagementPromptsPage: prompts, count + paging; text is question_text.
def _prompt(i: int, **overrides) -> dict:
    prompt = {"id": f"7d1c0b8e-1111-4c4c-9a9a-{i:012d}", "question_text": f"Question {i}?", "category": "fun",
              "active_date": None, "is_active": False, "response_count": i, "created_by": MEMBER,
              "created_at": "2026-09-01T10:00:00Z"}
    prompt.update(overrides)
    return prompt


class PromptPagingTest(ConsoleCaseTest):
    def test_prompts_page_with_filters_and_question_text(self):
        """Prompts are paged by Go; category/active/search reach it and cards show question_text. [case:console.engagement.engagement_prompts.renders]"""
        api = self.bff()
        api.list_engagement_prompts.return_value = _page("prompts", [_prompt(i) for i in range(25, 50)], 130, offset=25,
                                                         count=25)
        response = self.client.get(reverse("engagement_prompts"), {"page": "2", "category": "fun", "active": "no", "q": "question"})
        api.list_engagement_prompts.assert_called_once_with(limit=25, offset=25, category="fun", active="no", q="question",
                                                            sort="created_at", order="desc")
        self.assertContains(response, "Showing <strong>26–50</strong> of <strong>130</strong>")
        self.assertContains(response, "Question 25?")
        self.assertContains(response, reverse("engagement_prompt_activate", args=[_prompt(25)["id"]]))
        self.assertContains(response, 'class="list-toolbar"')

    def test_prompt_edit_finds_a_prompt_past_the_first_page(self):
        """Edit pages through prompts (Go returns 100 by default) to find the one asked for. [case:console.engagement.engagement_prompt_edit.performs]"""
        api = self.bff()
        late = _prompt(600, question_text="A late question?")
        api.list_engagement_prompts.side_effect = [_page("prompts", [_prompt(i) for i in range(500)], 501, limit=500),
                                                   _page("prompts", [late], 501, limit=500, offset=500)]
        response = self.client.get(reverse("engagement_prompt_edit", args=[late["id"]]))
        self.assertContains(response, "A late question?")
        self.assertEqual([c.kwargs for c in api.list_engagement_prompts.call_args_list],
                         [{"limit": 500, "offset": 0}, {"limit": 500, "offset": 500}])


# adminListEngagementNudges: nudges, count, clicked, by_type (page level), enabled + paging.
class NudgePagingTest(ConsoleCaseTest):
    def test_nudges_page_filters_and_page_level_counts(self):
        """Filters reach Go; the headline is Go's total and clicks are labelled as this page's. [case:console.engagement.engagement_nudges.renders]"""
        api = self.bff()
        nudges = [{"id": f"n{i}", "match_id": OTHER, "user_id": MEMBER, "counterparty_user_id": OTHER, "nudge_type": "stalled_24h",
                   "created_at": "2026-09-10T08:00:00Z", "clicked_at": "2026-09-10T09:00:00Z" if i < 3 else None} for i in range(10)]
        api.list_engagement_nudges.return_value = _page("nudges", nudges, 1234, limit=10, count=10, clicked=3,
                                                        by_type={"stalled_24h": 10}, enabled=False)
        response = self.client.get(reverse("engagement_nudges"), {"page_size": "10", "nudge_type": "stalled_24h", "clicked": "yes",
                                                                  "status": "clicked", "user_id": MEMBER, "from": "2026-09-01"})
        api.list_engagement_nudges.assert_called_once_with(limit=10, offset=0, nudge_type="stalled_24h", status="clicked",
                                                           clicked="yes", user_id=MEMBER, sort="created_at", order="desc",
                                                           **{"from": "2026-09-01"})
        self.assertEqual(response.context["nudge_count"], 1234)
        self.assertContains(response, "1,234")
        self.assertContains(response, "Clicked on this page")
        self.assertContains(response, ">Disabled<")
        self.assertContains(response, "stalled_24h · 10")
        self.assertContains(response, 'class="glass-table"')
        self.assertNotContains(response, 'class="table align-middle"')


# adminGroupCoversHandler: items (groupCoverReviewItem), count, status + paging.
class GroupCoverPagingTest(ConsoleCaseTest):
    module = "control_panel.views_group_covers"

    def test_queue_pages_and_searches_without_overriding_go_order(self):
        """Search and page reach Go; no sort is sent, so pending stays oldest first and approved newest first. [case:console.moderation_group_covers.group_covers.renders]"""
        api = self.bff()
        cover = {"cover_id": CASE, "group_id": OTHER, "group_name": "Sunday hikers", "group_kind": "community",
                 "owner_user_id": MEMBER, "uploaded_by": MEMBER, "status": "approved", "reason": "", "provider": "sandbox",
                 "mime_type": "image/jpeg", "width_px": 1200, "height_px": 675, "size_bytes": 2048,
                 "uploaded_at": "2026-10-01T09:00:00Z", "content_url": f"/v1/admin/moderation/group-covers/{CASE}/content"}
        api.group_covers.return_value = _page("items", [cover], 26, offset=25, count=1, status="approved")
        response = self.client.get(reverse("group_covers"), {"status": "approved", "q": "hikers", "page": "2"})
        api.group_covers.assert_called_once_with(limit=25, offset=25, status="approved", q="hikers")
        self.assertContains(response, "Showing <strong>26–26</strong> of <strong>26</strong>")
        self.assertContains(response, "Sunday hikers")
        self.assertContains(response, "approved, newest first")
        self.assertNotContains(response, 'href="?status=')

    def test_export_lists_covers_with_go_keys(self):
        """Excel carries the cover record (owner, check note, size) for the filtered queue. [case:console.lists.export.xlsx] [case:console.moderation_group_covers.group_covers.export]"""
        api = self.bff()
        api.group_covers.return_value = _page("items", [{"cover_id": CASE, "group_id": OTHER, "group_name": "Book club",
                                                         "owner_user_id": MEMBER, "status": "pending", "reason": "Needs a look",
                                                         "size_bytes": 4096, "uploaded_at": "2026-10-01T09:00:00Z"}], 1, limit=200)
        response = self.client.get(reverse("group_covers"), {"export": "xlsx"})
        rows, _ = _sheet(response)
        record = dict(zip(rows[0], rows[1]))
        self.assertEqual((record["Cover ID"], record["Group"], record["Owner ID"], record["Check note"], record["Bytes"]),
                         (CASE, "Book club", MEMBER, "Needs a look", 4096))
        self.assertEqual(api.group_covers.call_args.kwargs["status"], "pending")


# blogReviewHandler: cases, metrics, limit, offset, total, has_more.
def _blog_case(i: int, **overrides) -> dict:
    case = {"id": f"{i:08d}-0000-4000-8000-000000000000", "content_type": "post", "content_id": OTHER, "subject_id": MEMBER,
            "reason": "harassment", "description": f"Report {i}", "snapshot": {"title": "Private title", "body": "Private body"},
            "photo_ids": [], "status": "pending", "decision_note": "", "appeal": "", "version": 1,
            "created_at": "2026-09-30T10:00:00Z", "review_due_at": "2026-10-01T10:00:00Z", "overdue": i == 0,
            "evidence_purged": False}
    case.update(overrides)
    return case


class BlogReviewPagingTest(ConsoleCaseTest):
    module = "control_panel.views_blog"

    def test_queue_pages_filters_and_keeps_metrics_and_decisions(self):
        """Status, content type, search and page reach Go (due soonest first); metrics and decision forms stay. [case:console.moderation_blog.blog_reviews.renders] [case:console.lists.paging.server_side]"""
        api = self.bff()
        api.blog_reviews.return_value = APIResult(True, {"cases": [_blog_case(i) for i in range(10)], "metrics": {"pending_reviews": 41, "overdue_reviews": 2},
                                                         "limit": 10, "offset": 10, "total": 41, "has_more": True})
        response = self.client.get(reverse("blog_reviews"), {"status": "pending", "content_type": "post", "q": "harass",
                                                             "page": "2", "page_size": "10"})
        self.assertEqual(response.status_code, 200)
        api.blog_reviews.assert_called_once_with(limit=10, offset=10, status="pending", content_type="post", q="harass",
                                                 sort="review_due_at", order="asc")
        self.assertContains(response, "Showing <strong>11–20</strong> of <strong>41</strong>")
        self.assertContains(response, "pending_reviews")
        self.assertContains(response, "Record review decision", count=10)
        self.assertContains(response, f'action="{reverse("blog_decision", args=[_blog_case(0)["id"]])}"')
        self.assertIn("page=3", response.context["page"].next_url)
        self.assertNotContains(response, "offset=")
        self.assertIn("no-store", response["Cache-Control"])

    def test_sorting_by_report_date_descending_is_forwarded(self):
        """Choosing newest reports first reaches Go and stays in the page links. [case:console.moderation_blog.blog_reviews.renders]"""
        api = self.bff()
        api.blog_reviews.return_value = APIResult(True, {"cases": [_blog_case(i) for i in range(25)], "metrics": {}, "total": 30,
                                                         "limit": 25, "offset": 0})
        response = self.client.get(reverse("blog_reviews"), {"sort": "created_at", "dir": "desc", "status": "removed"})
        self.assertEqual(api.blog_reviews.call_args.kwargs, {"limit": 25, "offset": 0, "status": "removed", "sort": "created_at", "order": "desc"})
        self.assertIn("dir=desc", response.context["page"].next_url)

    def test_excel_export_has_every_case_but_no_reported_content(self):
        """Export pages through the whole queue; snapshots (private prose) are never written to the file. [case:console.lists.export.xlsx] [case:console.moderation_blog.blog_reviews.export]"""
        api = self.bff()
        api.blog_reviews.side_effect = [
            APIResult(True, {"cases": [_blog_case(i) for i in range(200)], "metrics": {}, "total": 260, "limit": 200, "offset": 0}),
            APIResult(True, {"cases": [_blog_case(i) for i in range(200, 260)], "metrics": {}, "total": 260, "limit": 200, "offset": 200}),
        ]
        response = self.client.get(reverse("blog_reviews"), {"export": "xlsx", "status": "dismissed"})
        self.assertEqual(response.status_code, 200)
        self.assertIn('filename="blog-review-cases-', response["Content-Disposition"])
        self.assertEqual([c.kwargs.get("offset") for c in api.blog_reviews.call_args_list], [0, 200])
        self.assertTrue(all(c.kwargs["status"] == "dismissed" for c in api.blog_reviews.call_args_list))
        rows, about = _sheet(response)
        self.assertEqual(len(rows), 261)
        record = dict(zip(rows[0], rows[1]))
        self.assertEqual(record["Case ID"], _blog_case(0)["id"])
        self.assertEqual(record["Report description"], "Report 0")
        self.assertEqual(record["Overdue"], "Yes")
        self.assertEqual(record["Review due (UTC)"], "2026-10-01T10:00:00Z")
        flat = repr(rows)
        self.assertNotIn("Private body", flat)
        self.assertNotIn("Private title", flat)
        self.assertEqual(about["Filter: status"], "dismissed")

    def test_upstream_failure_keeps_status_and_shows_filters(self):
        """A 403 from Go stays a 403 page with the banner and no decision forms. [case:console.moderation_blog.blog_reviews.renders]"""
        self.bff().blog_reviews.return_value = bff_error("Access denied", 403)
        response = self.client.get(reverse("blog_reviews"))
        self.assertContains(response, "Access denied", status_code=403)
        self.assertNotContains(response, "Record review decision", status_code=403)


# adminPhotoThemesHandler: themes (photoTheme) + paging.
class PhotoThemePagingTest(ConsoleCaseTest):
    module = "control_panel.views_photo_themes"

    def test_themes_page_in_display_order_with_status_filter(self):
        """Themes are paged by Go in display order; the status filter and search reach it. [case:console.engagement.photo_themes.renders]"""
        api = self.bff()
        themes = [{"id": f"t{i}", "slug": f"theme-{i}", "title": f"Theme {i}", "prompt": "Show us", "status": "archived",
                   "sort_order": i, "entry_count": 2, "my_entry_id": ""} for i in range(50)]
        api.photo_themes.return_value = _page("themes", themes, 75, limit=50)
        response = self.client.get(reverse("photo_themes"), {"status": "archived", "q": "theme"})
        api.photo_themes.assert_called_once_with(limit=50, offset=0, status="archived", q="theme", sort="sort_order", order="asc")
        self.assertContains(response, "Showing <strong>1–50</strong> of <strong>75</strong>")
        self.assertContains(response, "theme-49")
        self.assertContains(response, 'class="glass-table"')
        self.assertContains(response, reverse("photo_theme_save"))


# adminListFraudGraph: success, edges, automatic_enforcement + paging.
class GrowthFraudPagingTest(ConsoleCaseTest):
    def test_fraud_signals_are_listed_and_paged(self):
        """Fraud signals render from Go's edges; status, signal, member and sort reach Go; the metric is Go's total. [case:console.growth.growth_governance.renders]"""
        api = self.bff()
        api.growth_portfolio.return_value = APIResult(True, {"modules": []})
        edge = {"id": CASE, "left_member_id": MEMBER, "right_member_id": OTHER, "signal_type": "shared_device", "confidence": 0.91,
                "evidence": {"device_hash": "abc"}, "review_status": "open", "action_mode": "review_only", "reviewed_by": None,
                "reviewed_at": None, "created_at": "2026-09-20T10:00:00Z"}
        api.list_growth_fraud_graph.return_value = _page("edges", [edge], 140, success=True, automatic_enforcement=False)
        response = self.client.get(reverse("growth_governance"), {"signal_type": "shared_device", "member": MEMBER,
                                                                  "sort": "created_at", "dir": "asc"})
        api.list_growth_fraud_graph.assert_called_once_with(limit=25, offset=0, status="open", signal_type="shared_device",
                                                            member=MEMBER, sort="created_at", order="asc")
        self.assertContains(response, "140")
        self.assertContains(response, "shared device")
        self.assertContains(response, "device_hash: abc")
        self.assertContains(response, reverse("user_detail", args=[OTHER]))
        self.assertContains(response, 'class="list-toolbar"')

    def test_all_statuses_sends_no_status(self):
        """"All statuses" asks Go for every review status. [case:console.growth.growth_governance.renders]"""
        api = self.bff()
        api.list_growth_fraud_graph.return_value = _page("edges", [], 0)
        self.client.get(reverse("growth_governance"), {"status": "all"})
        self.assertEqual(api.list_growth_fraud_graph.call_args.kwargs["status"], "")


# adminListEconomyFraudCases: cases, count, status + paging.
class EconomyFraudPagingTest(ConsoleCaseTest):
    def test_fraud_cases_page_on_reconciliation(self):
        """Coin fraud cases are paged by Go with status, severity, member and sort; resolve forms stay; rules keep their own card. [case:console.billing.billing_reconciliation.renders]"""
        api = self.bff()
        case = {"id": CASE, "user_id": MEMBER, "rule_code": "gift_burst", "event_type": "gift_send", "status": "reviewing",
                "severity": "high", "recommended_action": "temporary_lock", "action_taken": "temporary_lock", "observed_value": 12,
                "trigger_value": 10, "window_seconds": 300, "occurrence_count": 2, "match_id": "", "receiver_user_id": OTHER,
                "first_detected_at": "2026-09-29T10:00:00Z", "last_detected_at": "2026-09-30T10:00:00Z",
                "lock_until": "2026-09-30T11:00:00Z", "resolution_note": ""}
        api.list_economy_fraud_cases.return_value = _page("cases", [case], 51, offset=50, count=1, status="reviewing")
        api.list_economy_fraud_rules.return_value = APIResult(True, {"rules": [{"rule_code": "gift_burst", "event_type": "gift_send"}]})
        response = self.client.get(reverse("billing_reconciliation"), {"status": "reviewing", "severity": "high", "user_id": MEMBER,
                                                                       "sort": "last_detected_at", "page": "3"})
        api.list_economy_fraud_cases.assert_called_once_with(limit=25, offset=50, status="reviewing", severity="high",
                                                             user_id=MEMBER, sort="last_detected_at", order="desc")
        self.assertContains(response, "Showing <strong>51–51</strong> of <strong>51</strong>")
        self.assertContains(response, "51 reviewing cases")
        self.assertContains(response, reverse("billing_fraud_case_resolve", args=[CASE]))
        self.assertContains(response, reverse("billing_fraud_rule_update", args=["gift_burst"]))
        # The window form keeps the fraud filters.
        self.assertContains(response, '<input type="hidden" name="severity" value="high">', html=False)

    def test_default_is_open_and_all_is_forwarded(self):
        """No status reads the open cases; "All statuses" sends all, which Go accepts. [case:console.billing.billing_reconciliation.renders]"""
        api = self.bff()
        api.list_economy_fraud_cases.return_value = _page("cases", [], 0)
        self.client.get(reverse("billing_reconciliation"))
        self.assertEqual(api.list_economy_fraud_cases.call_args.kwargs["status"], "open")
        self.client.get(reverse("billing_reconciliation"), {"status": "all"})
        self.assertEqual(api.list_economy_fraud_cases.call_args.kwargs["status"], "all")


# adminListProgressionFraud: cases, count + paging.
class ProgressionFraudPagingTest(ConsoleCaseTest):
    def test_xp_fraud_queue_pages_with_filters(self):
        """The XP fraud queue is paged by Go; status, severity, rule and search reach it; resolve forms stay. [case:console.progression.progression_admin.renders]"""
        api = self.bff()
        case = {"id": CASE, "user_id": MEMBER, "username": "member_one", "rule_code": "repeated_source_cap", "severity": "medium",
                "status": "reviewing", "evidence": {"rejected_attempts": 5}, "created_at": "2026-09-30T10:00:00Z"}
        api.list_progression_fraud.return_value = _page("cases", [case], 260, count=1)
        response = self.client.get(reverse("progression_admin"), {"status": "reviewing", "severity": "medium", "q": "member"})
        api.list_progression_fraud.assert_called_once_with(limit=25, offset=0, status="reviewing", severity="medium", q="member",
                                                           sort="created_at", order="desc")
        self.assertContains(response, "of <strong>260</strong>")
        self.assertContains(response, "@member_one")
        self.assertContains(response, reverse("progression_fraud_resolve", args=[CASE]))
        self.assertContains(response, 'class="glass-table"')
        self.assertNotContains(response, "table-dark")
        response = self.client.get(reverse("progression_admin"), {"status": "all"})
        self.assertEqual(api.list_progression_fraud.call_args.kwargs["status"], "")


class PagedListClientTest(ConsoleCaseTest):
    @patch("control_panel.services.go_client.GoBFFClient._request")
    def test_client_sends_only_contract_and_endpoint_filters(self, request):
        """Each list method sends limit/offset, the contract keys and its own filters, nothing else. [case:console.lists.params.allow_list]"""
        client = GoBFFClient(access_token="a", refresh_token="r", use_operator_context=False)
        client.list_catalog_gifts(limit=10, offset=20, category="roses", q="x", sort="name", order="asc", bogus="1")
        request.assert_called_with("GET", "/admin/catalog/gifts", params={"limit": 10, "offset": 20, "category": "roses", "q": "x", "sort": "name", "order": "asc"})
        client.blog_reviews(status="removed", limit=25, content_type="post", bogus="1")
        request.assert_called_with("GET", "/admin/moderation/blog", params={"status": "removed", "limit": 25, "content_type": "post"})
        client.list_engagement_nudges(limit=10, clicked="yes", match_id=OTHER)
        request.assert_called_with("GET", "/admin/engagement/nudges", params={"limit": 10, "clicked": "yes", "match_id": OTHER})
        client.list_growth_fraud_graph(status="", member=MEMBER)
        request.assert_called_with("GET", "/admin/growth/fraud-graph", params={"limit": 100, "member": MEMBER})
        client.list_economy_fraud_cases(status="", severity="high", offset=100)
        request.assert_called_with("GET", "/admin/billing/fraud/cases", params={"status": "open", "limit": 100, "severity": "high", "offset": 100})
        client.list_progression_fraud(status="open", rule_code="repeated_source_cap")
        request.assert_called_with("GET", "/admin/progression/fraud", params={"limit": 200, "status": "open", "rule_code": "repeated_source_cap"})
        client.photo_themes(limit=50, status="active")
        request.assert_called_with("GET", "/admin/engagement/photo-themes", params={"limit": 50, "status": "active"})
        client.list_engagement_prompts(limit=25, category="fun", active="yes")
        request.assert_called_with("GET", "/admin/engagement/prompts", params={"limit": 25, "category": "fun", "active": "yes"})
