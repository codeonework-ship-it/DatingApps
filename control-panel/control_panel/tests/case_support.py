"""Shared helpers for the catalog-case tests (qa/catalog/feature_catalog.json).

Each test method that proves a catalog case names it in its docstring as
``[case:<id>]`` for the QA Lab runner. The helpers here keep those tests
short: a logged-in operator (optionally with stored roles), a patched BFF
client whose every call answers ``APIResult(ok=True, data={})`` unless the
test says otherwise, and one ``assert_action_authz`` that checks the four
ways a console change must be refused.
"""
from __future__ import annotations

import base64
import csv
import io
import re
import zlib
from unittest.mock import MagicMock, patch

from django.contrib.messages import get_messages
from django.test import Client, TestCase

from control_panel.services.go_client import APIResult

OK = APIResult(True, {})


def bff_error(message="The service couldn't complete that request.", status=502):
    return APIResult(False, {}, message, status)


class AutoOKClient(MagicMock):
    """A GoBFFClient stand-in: any method not configured by the test answers
    ``APIResult(ok=True, data={})`` so a page renders its empty state."""

    def _get_child_mock(self, **kwargs):
        child = MagicMock(**kwargs)
        child.return_value = APIResult(True, {})
        return child


def login(client: Client, *, roles=None, username="console_admin", user_id="00000000-0000-0000-0000-0000000000aa"):
    session = client.session
    session["operator_access_token"] = "access-token"
    session["operator_refresh_token"] = "refresh-token"
    session["operator_username"] = username
    session["operator_user_id"] = user_id
    if roles is not None:
        session["operator_roles"] = list(roles)
    session.save()
    return client


def workbook(response) -> tuple[dict[str, list[tuple]], dict]:
    """Every sheet's rows (values only) and the About sheet as a dict."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(response.content))
    sheets = {name: list(wb[name].iter_rows(values_only=True)) for name in wb.sheetnames}
    about = {r[0]: r[1] for r in sheets.get("About", []) if r and r[0] is not None}
    return sheets, about


def csv_rows(response) -> list[list[str]]:
    """A CSV export's rows (the UTF-8 BOM is checked and dropped)."""
    assert response.content.startswith("\ufeff".encode()), "CSV must start with a UTF-8 BOM"
    return list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))


def pdf_text(response) -> str:
    """The text drawn on a ReportLab PDF's pages (ASCII85 + Flate streams
    decoded, every Tj string joined), so a test can assert what was printed."""
    raw = response.content
    assert raw.startswith(b"%PDF"), "not a PDF"
    parts = []
    for mo in re.finditer(rb"stream\r?\n(.*?)endstream", raw, re.S):
        data = mo.group(1).strip()
        try:
            if data.endswith(b"~>"):
                data = base64.a85decode(b"<~" + data if not data.startswith(b"<~") else data, adobe=True)
            data = zlib.decompress(data)
        except (ValueError, zlib.error):
            continue
        for text in re.findall(rb"\(((?:\\.|[^\\)])*)\)\s*Tj", data):
            # PDF string escapes: \ddd octal (WinAnsi, e.g. \267 is "·") or an escaped character.
            text = re.sub(rb"\\([0-7]{1,3})", lambda m: bytes([int(m.group(1), 8)]), text)
            parts.append(re.sub(rb"\\(.)", rb"\1", text).decode("cp1252", errors="replace"))
    return "".join(parts)


def flash(response) -> list[str]:
    return [str(m) for m in get_messages(response.wsgi_request)]


class ConsoleCaseTest(TestCase):
    """Base class: ``self.client`` is a logged-in operator with unknown roles."""

    module = "control_panel.views"

    def setUp(self):
        login(self.client)

    def bff(self, module: str | None = None) -> AutoOKClient:
        """Patch ``<module>.GoBFFClient`` (and views.GoBFFClient, used for the
        health badge) once per test and return the shared client instance."""
        existing = getattr(self, "_bff_api", None)
        if existing is not None:
            return existing
        api = AutoOKClient()
        for target in {module or self.module, "control_panel.views"}:
            patcher = patch(f"{target}.GoBFFClient")
            cls = patcher.start()
            cls.return_value = api
            self.addCleanup(patcher.stop)
        self._bff_api = api
        return api

    def as_roles(self, *roles) -> Client:
        return login(Client(), roles=list(roles))

    def assertFlash(self, response, text):
        messages = flash(response)
        self.assertTrue(any(text in m for m in messages), f"{text!r} not in {messages}")

    def assert_action_authz(
        self,
        url: str,
        data: dict,
        bff_method: str,
        *,
        refused_roles: tuple,
        allowed_roles: tuple = ("admin",),
        module: str | None = None,
        get_status: int = 405,
    ) -> None:
        """A console change is refused when it must be, and Go is never asked:

        * anonymous -> redirect to the login page;
        * GET on the action -> ``get_status`` (405 for POST-only actions, or
          200 for a form page) and no change;
        * POST without a CSRF token -> 403;
        * an operator whose stored roles Go would refuse -> redirected with
          an explanation;
        and an operator whose role Go allows does reach the BFF (so the gate
        is not simply refusing everything).
        """
        api = self.bff(module)
        method = getattr(api, bff_method)

        anonymous = Client().post(url, data)
        self.assertEqual(anonymous.status_code, 302, "anonymous POST must redirect")
        self.assertTrue(anonymous["Location"].startswith("/login/?next="), anonymous["Location"])
        method.assert_not_called()

        page = self.client.get(url)
        self.assertEqual(page.status_code, get_status, "GET on the action")
        method.assert_not_called()

        csrf_client = login(Client(enforce_csrf_checks=True), roles=["admin"])
        self.assertEqual(csrf_client.post(url, data).status_code, 403, "POST without a CSRF token")
        method.assert_not_called()

        for role in refused_roles:
            response = login(Client(), roles=[role]).post(url, data)
            self.assertEqual(response.status_code, 302, f"{role} must be redirected back")
            self.assertTrue(any("cannot make this change" in m for m in flash(response)), f"{role}: {flash(response)}")
            method.assert_not_called()

        for role in allowed_roles:
            method.reset_mock()
            login(Client(), roles=[role]).post(url, data)
            self.assertTrue(method.called, f"{role} should reach the BFF")
