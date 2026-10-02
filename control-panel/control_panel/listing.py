"""Server-side lists: one way to read paging, search, sort and filters from a
request, page through the Go BFF, render pagination, and export the whole
filtered result to Excel.

A list view declares a ``ListSpec`` (allowed filters, sorts, page sizes,
export columns) and a fetch function ``fetch(query, limit, offset) ->
(rows, total, error)`` that calls Go with exactly those parameters. Go does
the filtering, sorting and counting; the console never filters a page in
Python, so what the operator sees, the total, and the export always agree.

Unknown or invalid parameters are dropped, never forwarded: only values
listed in the spec reach Go.
"""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timezone
from typing import Any, Callable, Iterable
from urllib.parse import urlencode

from django.http import HttpRequest, HttpResponse
from django.shortcuts import render
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

PAGE_SIZES = (10, 25, 50, 100)
EXPORT_MAX_ROWS = 50_000
EXPORT_PAGE = 200
_DATE = re.compile(r"^\d{4}-\d{2}-\d{2}$")
# Excel would evaluate these as formulas when a cell starts with them.
_FORMULA_PREFIX = ("=", "+", "-", "@", "\t", "\r")


@dataclass(frozen=True)
class Filter:
    name: str
    label: str
    choices: tuple[tuple[str, str], ...] = ()  # (value, label); empty = free text
    kind: str = "choice"  # choice | text | date
    max_length: int = 100


@dataclass(frozen=True)
class Column:
    key: str
    label: str
    value: Callable[[dict], Any] | None = None  # defaults to row[key]
    width: int = 18


@dataclass(frozen=True)
class ListSpec:
    name: str  # used for the export file name
    filters: tuple[Filter, ...] = ()
    sorts: tuple[tuple[str, str], ...] = ()  # (value, label), first is default
    default_page_size: int = 25
    searchable: bool = True
    search_label: str = "Search"
    columns: tuple[Column, ...] = ()  # Excel export


@dataclass
class ListQuery:
    spec: ListSpec
    page: int = 1
    page_size: int = 25
    q: str = ""
    sort: str = ""
    direction: str = "desc"
    filters: dict[str, str] = field(default_factory=dict)

    @property
    def offset(self) -> int:
        return (self.page - 1) * self.page_size

    def params(self, **overrides: Any) -> dict[str, str]:
        """The query string of this list (only non-default values)."""
        values: dict[str, Any] = {
            "q": self.q,
            **self.filters,
            "sort": self.sort if self.sort != self._default_sort else "",
            "dir": self.direction if self.direction != "desc" else "",
            "page_size": self.page_size if self.page_size != self.spec.default_page_size else "",
            "page": self.page if self.page != 1 else "",
        }
        values.update(overrides)
        return {k: str(v) for k, v in values.items() if v not in ("", None)}

    def url(self, **overrides: Any) -> str:
        params = self.params(**overrides)
        return "?" + urlencode(params) if params else "?"

    def go_params(self) -> dict[str, Any]:
        """Search, filters and sort as Go reads them (paging is separate)."""
        out: dict[str, Any] = {k: v for k, v in self.filters.items() if v}
        if self.q:
            out["q"] = self.q
        if self.sort:
            out["sort"] = self.sort
            out["order"] = self.direction
        return out

    @property
    def _default_sort(self) -> str:
        return self.spec.sorts[0][0] if self.spec.sorts else ""

    @property
    def active_filter_count(self) -> int:
        return sum(1 for v in self.filters.values() if v) + (1 if self.q else 0)


def _int(value: Any, default: int) -> int:
    try:
        return int(str(value).strip())
    except (TypeError, ValueError):
        return default


def parse(request: HttpRequest, spec: ListSpec) -> ListQuery:
    g = request.GET
    page_size = _int(g.get("page_size"), spec.default_page_size)
    if page_size not in PAGE_SIZES:
        page_size = spec.default_page_size
    sorts = {value for value, _ in spec.sorts}
    sort = g.get("sort", "").strip()
    if sort not in sorts:
        sort = spec.sorts[0][0] if spec.sorts else ""
    direction = "asc" if g.get("dir", "").strip().lower() == "asc" else "desc"
    filters: dict[str, str] = {}
    for f in spec.filters:
        raw = g.get(f.name, "").strip()[: f.max_length]
        if f.kind == "choice" and f.choices and raw not in {v for v, _ in f.choices}:
            raw = ""
        if f.kind == "date" and raw and not _DATE.match(raw):
            raw = ""
        filters[f.name] = raw
    q = g.get("q", "").strip()[:200] if spec.searchable else ""
    return ListQuery(spec=spec, page=max(1, _int(g.get("page"), 1)), page_size=page_size,
                     q=q, sort=sort, direction=direction, filters=filters)


Fetch = Callable[[ListQuery, int, int], "tuple[list[dict], int | None, str]"]


@dataclass
class Page:
    query: ListQuery
    rows: list[dict]
    total: int | None  # None when Go does not count
    error: str = ""

    @property
    def pages(self) -> int | None:
        if self.total is None:
            return None
        return max(1, -(-self.total // self.query.page_size))

    @property
    def start(self) -> int:
        return self.query.offset + 1 if self.rows else 0

    @property
    def end(self) -> int:
        return self.query.offset + len(self.rows)

    @property
    def has_previous(self) -> bool:
        return self.query.page > 1

    @property
    def has_next(self) -> bool:
        if self.total is not None:
            return self.end < self.total
        return len(self.rows) >= self.query.page_size

    @property
    def previous_url(self) -> str:
        return self.query.url(page=self.query.page - 1 if self.query.page > 2 else "")

    @property
    def next_url(self) -> str:
        return self.query.url(page=self.query.page + 1)

    def links(self) -> list[dict]:
        """Numbered page links with gaps: 1 … 4 5 [6] 7 8 … 20."""
        pages = self.pages
        current = self.query.page
        if pages is None:
            return []
        wanted = {1, pages, *range(current - 2, current + 3)}
        out: list[dict] = []
        last = 0
        for number in sorted(n for n in wanted if 1 <= n <= pages):
            if number - last > 1:
                out.append({"gap": True})
            out.append({"number": number, "url": self.query.url(page=number if number > 1 else ""), "current": number == current})
            last = number
        return out

    def sort_link(self, key: str) -> dict:
        """Header link that sorts by ``key`` (toggles direction if active)."""
        active = self.query.sort == key
        direction = "asc" if active and self.query.direction == "desc" else "desc"
        return {"url": self.query.url(sort=key, dir=direction, page=""), "active": active,
                "direction": self.query.direction if active else "",
                "aria": "descending" if active and self.query.direction == "desc" else ("ascending" if active else "none")}


def load(query: ListQuery, fetch: Fetch) -> Page:
    rows, total, error = fetch(query, query.page_size, query.offset)
    page = Page(query=query, rows=rows, total=total, error=error)
    # A page past the end (e.g. after filtering) falls back to the last page.
    if not rows and not error and total and query.page > 1 and page.pages and query.page > page.pages:
        query.page = page.pages
        rows, total, error = fetch(query, query.page_size, query.offset)
        page = Page(query=query, rows=rows, total=total, error=error)
    return page


def context(page: Page) -> dict:
    """Template context for ``_list_controls.html`` and ``_pagination.html``."""
    q = page.query
    return {
        "page": page,
        "list_query": q,
        "list_spec": q.spec,
        "page_sizes": PAGE_SIZES,
        "sorts": [{"value": v, "label": label, "selected": v == q.sort} for v, label in q.spec.sorts],
        "list_filters": [{"spec": f, "value": q.filters.get(f.name, "")} for f in q.spec.filters],
        "export_url": q.url(export="xlsx", page="", page_size=""),
        "reset_url": "?",
        "sort_links": {key: page.sort_link(key) for key, _ in q.spec.sorts},
    }


# ── Excel export ──────────────────────────────────────────────────────────

def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (datetime, date)):
        return value.replace(tzinfo=None) if isinstance(value, datetime) else value
    if isinstance(value, (list, tuple, set)):
        value = ", ".join(str(v) for v in value)
    elif isinstance(value, dict):
        value = ", ".join(f"{k}: {v}" for k, v in value.items())
    text = str(value)
    if text.startswith(_FORMULA_PREFIX):
        text = "'" + text  # never a formula (CSV/Excel injection)
    return text[:32_000]


def export_rows(query: ListQuery, fetch: Fetch) -> tuple[Iterable[dict], str]:
    """Every row matching the query, paged through Go; stops at the cap."""
    collected: list[dict] = []
    offset = 0
    while len(collected) < EXPORT_MAX_ROWS:
        rows, total, error = fetch(query, EXPORT_PAGE, offset)
        if error:
            return collected, error
        collected.extend(rows)
        offset += len(rows)
        if len(rows) < EXPORT_PAGE or (total is not None and offset >= total):
            break
    return collected[:EXPORT_MAX_ROWS], ""


def xlsx_response(query: ListQuery, fetch: Fetch, *, title: str) -> HttpResponse:
    rows, error = export_rows(query, fetch)
    if error and not rows:
        return HttpResponse(f"Export failed: {error}", status=502, content_type="text/plain; charset=utf-8")
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title[:31] or "Export")
    columns = query.spec.columns
    for i, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = column.width
    ws.freeze_panes = "A2"
    header = []
    for column in columns:
        cell = WriteOnlyCell(ws, value=column.label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2D2A5A")
        header.append(cell)
    ws.append(header)
    for row in rows:
        ws.append([_cell(c.value(row) if c.value else row.get(c.key)) for c in columns])
    # A second sheet records what was exported, so a file is self-describing.
    about = wb.create_sheet("About")
    about.append(["Report", title])
    about.append(["Exported (UTC)", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S")])
    about.append(["Rows", len(rows)])
    for key, value in query.params(page="", page_size="").items():
        about.append([f"Filter: {key}", value])
    if len(rows) >= EXPORT_MAX_ROWS:
        about.append(["Note", f"Capped at {EXPORT_MAX_ROWS:,} rows; narrow the filters for the rest."])
    if error:
        about.append(["Note", f"Stopped early: {error}"])
    buffer = io.BytesIO()
    wb.save(buffer)
    stamp = datetime.now(timezone.utc).strftime("%Y%m%d-%H%M")
    response = HttpResponse(buffer.getvalue(),
                            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{query.spec.name}-{stamp}.xlsx"'
    response["Cache-Control"] = "no-store"
    return response


def respond(request: HttpRequest, spec: ListSpec, fetch: Fetch, *, title: str,
            render_page: Callable[[Page], HttpResponse]) -> HttpResponse:
    """The whole list flow: Excel when ``?export=xlsx``, else one page."""
    query = parse(request, spec)
    if request.GET.get("export") == "xlsx" and spec.columns:
        return xlsx_response(query, fetch, title=title)
    return render_page(load(query, fetch))


def simple_view(request: HttpRequest, spec: ListSpec, call: Callable[..., Any], *, items_key: str, template: str,
                title: str, base_context: Callable[[], dict], context_name: str,
                map_filters: Callable[[ListQuery], dict[str, Any]] | None = None) -> HttpResponse:
    """A list page whose Go method takes ``limit``, ``offset``, its filters
    and the contract extras (q, sort, order, from, to)."""
    def fetch(query: ListQuery, limit: int, offset: int):
        kwargs = map_filters(query) if map_filters else query.go_params()
        result = call(limit=limit, offset=offset, **kwargs)
        if not result.ok:
            return [], None, result.error
        data = result.data if isinstance(result.data, dict) else {}
        rows = [r for r in data.get(items_key) or [] if isinstance(r, dict)]
        total = data.get("total")
        return rows, total if isinstance(total, int) and not isinstance(total, bool) else None, ""

    def render_page(page: Page) -> HttpResponse:
        ctx = base_context()
        ctx.update(context(page))
        ctx.update({context_name: page.rows, "total": page.total, "error": page.error})
        return render(request, template, ctx)

    return respond(request, spec, fetch, title=title, render_page=render_page)
